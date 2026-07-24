from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ACCENT = RGBColor(135, 0, 18)
TEXT = RGBColor(0, 16, 36)
MUTED = RGBColor(74, 85, 104)


def set_cell_text_font(run, size=8.2, bold=False, color=TEXT):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_paragraph_spacing(p, before=0, after=0, line=0.95):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def set_run(run, size=8.2, bold=False, color=TEXT):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_section(doc: Document, title: str):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=4, after=2, line=1.0)
    r = p.add_run("■ " + title)
    set_run(r, size=11, bold=True, color=ACCENT)
    return p


def add_role(doc: Document, title: str, date: str | None = None):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(14.3)
    table.columns[1].width = Cm(3.1)
    for cell in table.rows[0].cells:
        cell.margin_top = 0
        cell.margin_bottom = 0
        cell.margin_left = 0
        cell.margin_right = 0
    left = table.cell(0, 0).paragraphs[0]
    set_paragraph_spacing(left, before=0, after=0, line=1.0)
    r = left.add_run(title)
    set_run(r, size=9.1, bold=True)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(right, before=0, after=0, line=1.0)
    if date:
        rr = right.add_run(date)
        set_run(rr, size=8.8, color=MUTED)
    return table


def add_inline(doc: Document, parts):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=0, line=0.95)
    for text, bold in parts:
        r = p.add_run(text)
        set_run(r, size=8.15, bold=bold)


def add_bullet(doc: Document, label: str, text: str):
    p = doc.add_paragraph(style=None)
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.first_line_indent = Cm(-0.18)
    set_paragraph_spacing(p, before=0, after=1, line=0.93)
    bullet = p.add_run("· ")
    set_run(bullet, size=8.0, color=ACCENT, bold=True)
    r = p.add_run(label)
    set_run(r, size=8.0, bold=True)
    r2 = p.add_run(text)
    set_run(r2, size=8.0)


def strip_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tblPr.append(borders)


def build(out: Path, omit_exchange: bool):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.05)
    sec.bottom_margin = Cm(1.05)
    sec.left_margin = Cm(1.15)
    sec.right_margin = Cm(1.15)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(8.1)
    normal.font.color.rgb = TEXT

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=1, line=1)
    r = p.add_run("李安")
    set_run(r, size=17, bold=True, color=TEXT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=3, line=1)
    r = p.add_run("18810917395 ｜ ulnachat@gmail.com ｜ 求职意向：互联网产品经理 / AI 产品 / 内容产品")
    set_run(r, size=8.5, color=MUTED)

    add_section(doc, "教育经历")
    add_role(doc, "芝加哥大学｜法学院｜法学硕士 LL.M.", "预计 2026.09 入学")
    add_role(doc, "北京大学｜法学院｜国际法学｜硕士", "2023.08–2025.07")
    if omit_exchange:
        add_inline(doc, [("GPA：", True), ("3.72/4.00", False)])
    else:
        add_inline(doc, [("GPA：", True), ("3.72/4.00；", False), ("国际交流经历：", True), ("日本东京大学法学院交换项目、荷兰海牙国际法学院冬季课程", False)])
    add_role(doc, "北京大学｜法学院｜法学｜本科", "2019.08–2023.07")
    if omit_exchange:
        add_inline(doc, [("GPA：", True), ("3.58/4.00（前 30%），保研本院；", False), ("奖励情况：", True), ("北京大学新生奖学金、三等奖学金（20/182）、学习优秀奖", False)])
    else:
        add_inline(doc, [("GPA：", True), ("3.58/4.00（前 30%），保研本院；", False), ("奖励情况：", True), ("北京大学新生奖学金、三等奖学金（20/182）、学习优秀奖；", False), ("国际交流经历：", True), ("法国巴黎政治学院交换项目", False)])

    add_section(doc, "项目经历")
    add_role(doc, "AnimeLore —— 二次元 AI 内容图鉴｜个人独立项目", "2026.04–至今")
    add_inline(doc, [("", False), ("面向二次元深度用户的 AI 作品图鉴，整合 AniList、Wikipedia、萌娘百科等多源数据，生成结构化作品解读与个性化探索体验。", False)])
    add_bullet(doc, "多源 RAG 聚合 + 四维解读卡：", "设计“硬信息优先、社区语境补充”的数据融合策略，整合 AniList/Wikipedia 可核验事实与萌娘百科、NGA 等中文社区梗/考据，并引入第二 AI 交叉验证机制辅助校验关键事实，确保内容准确且有温度。")
    add_bullet(doc, "冷知识系统与知识漫游：", "识别到用户具有广泛探索二次元文化的需求，设计预生成精品库 + 加权轮换机制，以卡片流形式提供低门槛、可自由漫游的文化探索入口，覆盖梗起源、制作幕后、概念沿革等内容类型。")
    add_bullet(doc, "芽衣 AI 对话 + 社交分享闭环：", "发现动漫观看是强情绪体验，设计定制人格的 AI 聊天助手（二次元同好视角），并在用户聊出实质感想后触发分享卡片生成，支持导出到小红书等平台，将个人的知识消费转化为社交内容。")
    add_bullet(doc, "技术实现：", "使用 Next.js 14 + TypeScript + DeepSeek API 独立开发，零编程经验背景下通过 AI 辅助编程完成全栈实现，包含 SSE 流式渲染、自适应 RAG 管线、html2canvas 图片生成等能力。")

    add_section(doc, "实习经历")
    add_role(doc, "霍金路伟律师事务所｜数据合规｜实习生", "2023.03–2023.05")
    add_bullet(doc, "数据跨境规则研究：", "围绕数据跨境流动、数据本地化及个人信息出境议题，梳理 RCEP、CPTPP、DEPA、WTO 电子商务谈判、APEC CBPR、GDPR 等十余个国际规则框架，逐项分析数据自由流动、计算设施本地化、个人隐私保护、国家安全例外等条款，形成完整调研报告。")
    add_bullet(doc, "尽调与信息支持：", "参与数据合规法规、合同及配套材料的英文翻译与校对；通过公开渠道检索企业经营、政策环境及行业监管相关信息，支持尽职调查、客户公共事务分析及合规判断。")
    add_role(doc, "汉坤律师事务所｜争议解决组｜实习生", "2021.07–2021.08")
    add_bullet(doc, "案件信息结构化：", "支持多类型商事争议工作，承担法律检索、证据材料整理、合同翻译校对、企业涉诉及工商信息检索等任务；围绕外国判决承认与执行、委托贷款监管处罚、刑民交叉、股权历史变化等议题，形成案例梳理表、监管处罚表和股权变化表等结构化材料。")

    add_section(doc, "研究与活动经历")
    add_role(doc, "北京大学国际法研究所｜助教、研究助理", "2023.03–2025.07")
    add_bullet(doc, "课程与内容运营：", "两次担任全英文国际法课程助教，单门课程服务超过 200 名学生，负责课程宣传、报名统计、材料管理、课堂展示排期与师生沟通；负责研究所公众号讲座推送、演讲稿整理及 40 周年校庆 H5/推送产出，单篇阅读量达数千。")
    add_role(doc, "北京大学“挑战杯”学术科创竞赛｜队长｜一等奖", "2020.09–2021.05")
    add_bullet(doc, "研究设计与项目推进：", "带领 3 人团队围绕新冠疫情初期公共决策开展法律经济学研究，结合流行病学与经济学模型进行成本收益分析；负责研究框架搭建、英文文献阅读、团队分工及 40 余页论文/参赛材料撰写，作品获法学院院内第一名、北京大学“挑战杯”一等奖。")
    add_role(doc, "北京大学思维潜能开发协会｜副会长", "2022.10–2025.07")
    add_bullet(doc, "社群运营与合作拓展：", "参与管理 3000 余人剧本杀玩家社群及 100 余名学生主持人，负责主持人培训、迎新宣传、公众号运营、社团规则制定与活动流程优化；邀请知名剧本杀作者入校分享，并与学校周边店家谈判价格折扣与场地合作。")

    add_section(doc, "其他")
    add_bullet(doc, "语言能力：", "英语 TOEFL 114，CET-6 612；日语 JLPT N1，可进行英文/日文资料检索、阅读与整理。")
    add_bullet(doc, "AI 工具与产品实践：", "熟练使用 ChatGPT、Claude、DeepSeek 等大模型辅助研究、写作、翻译、信息整理与需求拆解；可通过 Codex、Cursor、Antigravity 等 AI coding 工具进行基础前后端开发、页面搭建与迭代，并结合 Stitch、Claude Design 等工具完成前端原型与视觉探索；具备基于个人使用场景进行需求拆解和产品构思的实践，熟悉资料调研、结构化文档写作、内容策略设计、Prompt 设计与 AI 生成结果评估。")

    for table in doc.tables:
        strip_table_borders(table)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True)
    parser.add_argument("--omit-exchange", action="store_true")
    args = parser.parse_args()
    build(Path(args.out), args.omit_exchange)


if __name__ == "__main__":
    main()
